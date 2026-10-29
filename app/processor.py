"""
生成对账 Excel + 清理临时文件
"""
from __future__ import annotations

import io
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from .reconciler import run_reconcile


def _header(ws, row=1):
    fill = PatternFill("solid", fgColor="1E3A5F")
    font = Font(bold=True, color="FFFFFF")
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _width(ws, max_w=36):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 10
        for cell in col:
            try:
                length = max(length, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[letter].width = min(length + 2, max_w)


def build_excel(result: dict, metas: list[dict]) -> bytes:
    wb = Workbook()

    # 汇总
    ws = wb.active
    ws.title = "对账汇总"
    ws.append(["项目", "数值"])
    _header(ws)
    summary = result.get("summary") or {}
    for k, v in summary.items():
        if isinstance(v, dict):
            ws.append([k, str(v)])
        else:
            ws.append([k, v])
    ws.append([])
    ws.append(["【解析信息】", ""])
    for m in metas:
        ws.append([
            m.get("filename"),
            f"{m.get('tag')} | {m.get('platform')} | 有效{m.get('rows_valid')}行",
        ])
        for w in m.get("warnings") or []:
            ws.append(["提示", w])
    _width(ws)

    # 匹配成功
    ws2 = wb.create_sheet("匹配成功")
    matched = result.get("matched")
    if matched is not None and not matched.empty:
        for r in dataframe_to_rows(matched, index=False, header=True):
            ws2.append(r)
        _header(ws2)
        _width(ws2)
    else:
        ws2.append(["无跨文件匹配记录（可能只有单文件，或订单号/日期对不上）"])

    # 金额不一致
    ws3 = wb.create_sheet("金额不一致")
    mm = result.get("amount_mismatch")
    if mm is not None and not mm.empty:
        for r in dataframe_to_rows(mm, index=False, header=True):
            ws3.append(r)
        _header(ws3)
        red = PatternFill("solid", fgColor="FFC7CE")
        for row in ws3.iter_rows(min_row=2):
            for c in row:
                c.fill = red
        _width(ws3)
    else:
        ws3.append(["无金额不一致项"])
        ws3["A1"].font = Font(color="006600")

    # 仅在各文件
    only = result.get("only") or {}
    for label, df in only.items():
        safe_title = f"仅在{label}"[:28]
        wsx = wb.create_sheet(safe_title)
        if df is not None and not df.empty:
            for r in dataframe_to_rows(df.head(5000), index=False, header=True):
                wsx.append(r)
            _header(wsx)
            _width(wsx)
        else:
            wsx.append(["无独有记录"])

    # 全部标准化
    ws_all = wb.create_sheet("全部标准化流水")
    all_df = result.get("all")
    if all_df is not None and not all_df.empty:
        for r in dataframe_to_rows(all_df.head(8000), index=False, header=True):
            ws_all.append(r)
        _header(ws_all)
        _width(ws_all)
    else:
        ws_all.append(["无数据"])

    # 隐私
    ws_p = wb.create_sheet("隐私说明")
    for line in [
        "本工具隐私承诺",
        "1. 上传文件仅临时处理，完成后立即删除。",
        "2. 不建立账号，不落库，不用于训练或第三方。",
        "3. 对账结果仅供参考，请人工复核后再用于结算/报税。",
        "4. 建议处理完立即下载并关闭页面。",
    ]:
        ws_p.append([line])
    ws_p.column_dimensions["A"].width = 72

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def process_uploads(paths: list[Path], names: list[str]) -> tuple[bytes, dict]:
    result, metas = run_reconcile(paths, names)
    if not any(not (result.get("all") is None or result["all"].empty) for _ in [0]):
        # still build empty-ish report
        pass
    if result.get("all") is None or (hasattr(result["all"], "empty") and result["all"].empty):
        # 尝试仍输出带警告的汇总
        if all((m.get("rows_valid") or 0) == 0 for m in metas):
            raise ValueError(
                "未能解析出有效账单数据。请确认文件含「订单号/金额」等列，格式为 Excel 或 CSV。"
            )

    excel = build_excel(result, metas)
    matched = result.get("matched")
    mismatch = result.get("amount_mismatch")
    summary = {
        "ok": True,
        "summary": result.get("summary"),
        "matched": 0 if matched is None or matched.empty else len(matched),
        "mismatch": 0 if mismatch is None or mismatch.empty else len(mismatch),
        "metas": metas,
    }
    return excel, summary


def safe_cleanup(paths: list[Path]):
    for p in paths:
        try:
            if p.exists():
                p.unlink()
        except Exception:
            pass
